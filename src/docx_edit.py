"""Edit caption paragraphs and lists while retaining original object anchors."""

from io import BytesIO

from pandas import isna

from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from src.caption_format import build_formatted_source, clean_caption_cell, clean_page
from src.object_detect import (
    detect_list_of_figures_content_range,
    detect_list_of_tables_content_range,
)


def _row_value(row, *names, default=""):
    for name in names:
        if name in row and not isna(row[name]):
            return row[name]

    return default


def build_caption_parts(row, caption_settings):
    caption_label = clean_caption_cell(
        _row_value(row, "caption_label", default=caption_settings["caption_label"])
    )
    new_number = clean_caption_cell(_row_value(row, "new_number", "number"))
    caption_text = clean_caption_cell(_row_value(row, "caption_text", "text"))
    source_part = build_formatted_source(_row_value(row, "source_text", "source"), caption_settings)
    label_part = f"{caption_label} {new_number}. " if new_number else f"{caption_label} "
    text_part = f"{caption_text}{' ' if source_part else ''}" if caption_text else ""

    return {
        "label_part": label_part,
        "text_part": text_part,
        "source_part": source_part,
    }


def build_caption_main_parts(row, caption_settings):
    parts = build_caption_parts(row, caption_settings)

    return {
        "label_part": parts["label_part"],
        "text_part": parts["text_part"],
        "source_part": "",
    }


def build_source_part(row, caption_settings):
    source_part = build_formatted_source(_row_value(row, "source_text", "source"), caption_settings)

    return {
        "label_part": "",
        "text_part": "",
        "source_part": source_part,
    }


def clear_paragraph(paragraph):
    paragraph_element = paragraph._element

    for child in list(paragraph_element):
        if child.tag != qn("w:pPr"):
            paragraph_element.remove(child)


def _set_run_style(run, caption_settings, part):
    run.font.name = caption_settings["font_name"]
    run.font.size = Pt(caption_settings["font_size"])
    run.bold = caption_settings[f"{part}_bold"]
    run.italic = caption_settings[f"{part}_italic"]


def write_caption_to_paragraph(paragraph, caption_parts, caption_settings):
    clear_paragraph(paragraph)

    run_specs = [
        ("label_part", "label"),
        ("text_part", "text"),
        ("source_part", "source"),
    ]

    for part_key, style_key in run_specs:
        text = caption_parts[part_key]

        if not text:
            continue

        run = paragraph.add_run(text)
        _set_run_style(run, caption_settings, style_key)


def _body_width(doc):
    section = doc.sections[-1]

    return section.page_width - section.left_margin - section.right_margin


def _set_table_of_figures_run_style(run, caption_settings):
    run.font.name = caption_settings["font_name"]
    run.font.size = Pt(caption_settings["font_size"])


def _remove_paragraph(paragraph):
    element = paragraph._element
    parent = element.getparent()
    parent.remove(element)


def _existing_paragraph(paragraphs, index):
    if isna(index):
        return None
    index = int(index)
    return paragraphs[index] if 0 <= index < len(paragraphs) else None


def _insert_paragraph(doc, anchor, before=False):
    element = OxmlElement("w:p")
    anchor.addprevious(element) if before else anchor.addnext(element)
    return Paragraph(element, doc._body)


def edit_docx_captions(doc, caption_rows, caption_settings):
    """Update existing captions and create missing paragraphs at object anchors."""
    rows = (
        [row for _, row in caption_rows.iterrows()]
        if hasattr(caption_rows, "iterrows") else list(caption_rows)
    )
    paragraphs, tables = list(doc.paragraphs), list(doc.tables)
    actions = []

    # Resolve every original XML anchor before any insertion changes paragraph indices.
    for row in rows:
        if not clean_caption_cell(_row_value(row, "caption_text", "text")):
            continue
        main = _existing_paragraph(
            paragraphs, _row_value(row, "caption_paragraph_idx", default=None)
        )
        object_type = _row_value(row, "object_type", default=None)
        table = None

        if main is None and object_type == "figure":
            object_idx = _row_value(row, "object_paragraph_idx", default=None)
            anchor = _existing_paragraph(paragraphs, object_idx)
            main = (anchor._element, False) if anchor is not None else None
        elif object_type == "table":
            table_idx = _row_value(row, "table_idx", default=None)
            if not isna(table_idx) and 0 <= int(table_idx) < len(tables):
                table = tables[int(table_idx)]
            if main is None and table is not None:
                position = _row_value(
                    row, "table_caption_position",
                    default=caption_settings.get("table_caption_position", "caption_above"),
                )
                main = (table._element, position != "caption_below")

        if main is None:
            continue

        separate_source = (
            object_type == "table"
            and _row_value(
                row, "table_caption_position",
                default=caption_settings.get("table_caption_position"),
            ) == "label_above_source_below"
        )
        source = None
        if separate_source:
            source = _existing_paragraph(
                paragraphs, _row_value(row, "source_paragraph_idx", default=None)
            )
            if source is None and table is not None and clean_caption_cell(
                _row_value(row, "source_text", "source")
            ):
                source = (table._element, False)
        actions.append((row, main, source, separate_source))

    for row, main, source, separate_source in actions:
        paragraph = (
            _insert_paragraph(doc, main[0], before=main[1])
            if isinstance(main, tuple) else main
        )
        if separate_source:
            write_caption_to_paragraph(
                paragraph, build_caption_main_parts(row, caption_settings), caption_settings
            )
            if source is not None:
                source_paragraph = (
                    _insert_paragraph(doc, source[0], before=source[1])
                    if isinstance(source, tuple) else source
                )
                write_caption_to_paragraph(
                    source_paragraph, build_source_part(row, caption_settings), caption_settings
                )
        else:
            write_caption_to_paragraph(
                paragraph, build_caption_parts(row, caption_settings), caption_settings
            )

    return doc


def _caption_list_elements(doc, rows_to_append, caption_settings, title):
    """Build a list at the document end and return its movable XML elements."""
    elements = []
    title_paragraph = doc.add_paragraph(title)
    title_paragraph.style = "Heading 1"
    elements.append(title_paragraph._element)

    for run in title_paragraph.runs:
        run.font.name = caption_settings["font_name"]

    if hasattr(rows_to_append, "iterrows"):
        rows = (row for _, row in rows_to_append.iterrows())
    else:
        rows = iter(rows_to_append)

    for row in rows:
        label = clean_caption_cell(_row_value(row, "label"))
        caption_text = clean_caption_cell(_row_value(row, "text", "caption_text"))
        page = clean_page(_row_value(row, "page"))

        if not label or not caption_text:
            continue

        paragraph = doc.add_paragraph()
        elements.append(paragraph._element)
        paragraph.paragraph_format.tab_stops.clear_all()
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            _body_width(doc),
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.DOTS,
        )

        label_run = paragraph.add_run(f"{label} ")
        label_run.bold = True
        _set_table_of_figures_run_style(label_run, caption_settings)

        text_run = paragraph.add_run(caption_text)
        _set_table_of_figures_run_style(text_run, caption_settings)

        if page:
            page_run = paragraph.add_run(f"\t{page}")
            _set_table_of_figures_run_style(page_run, caption_settings)

    return elements


def _append_caption_list(doc, rows_to_append, caption_settings, title):
    doc.add_paragraph()
    _caption_list_elements(doc, rows_to_append, caption_settings, title)
    return doc


def _list_title(caption_settings, list_type):
    polish = caption_settings.get("language") == "Polish"
    if list_type == "figures":
        return "Spis rysunków" if polish else "List of Figures"
    return "Lista tabel" if polish else "List of Tables"


def replace_caption_lists(
    doc,
    figure_rows,
    table_rows,
    caption_settings,
    replace_figures=True,
    replace_tables=True,
):
    """Replace selected lists at their original positions."""
    replacements = []
    if replace_tables:
        replacements.append((
            *detect_list_of_tables_content_range(doc),
            table_rows,
            _list_title(caption_settings, "tables"),
        ))
    if replace_figures:
        replacements.append((
            *detect_list_of_figures_content_range(doc),
            figure_rows,
            _list_title(caption_settings, "figures"),
        ))

    # Work backwards and insert before the old heading so each list keeps its location.
    paragraphs = list(doc.paragraphs)
    anchored = sorted(
        (item for item in replacements if item[0] is not None),
        key=lambda item: item[0], reverse=True,
    )
    for start, end, rows, title in anchored:
        anchor = paragraphs[start]._element
        if rows is not None and len(rows):
            for element in _caption_list_elements(doc, rows, caption_settings, title):
                anchor.addprevious(element)
        for paragraph in paragraphs[start:end]:
            _remove_paragraph(paragraph)

    # This fallback is used only when a selected list has no detectable heading.
    for start, _, rows, title in replacements:
        if start is None and rows is not None and len(rows):
            _append_caption_list(doc, rows, caption_settings, title)

    return doc


def docx_to_bytes(doc):
    buffer = BytesIO()
    doc.save(buffer)

    return buffer.getvalue()
