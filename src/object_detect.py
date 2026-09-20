import re
import pandas as pd
from docx.oxml.ns import qn

# Tables
def extract_tables(doc):
    """
    Parse tables from a .docx document and return a DataFrame.
    """
    rows = []
    table_counter = 0
    paragraph_counter = 0
    previous_paragraph_idx = None

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            previous_paragraph_idx = paragraph_counter
            paragraph_counter += 1

        elif child.tag == qn("w:tbl"):
            table_counter += 1

            rows.append({
                "table_no": table_counter,
                "table_idx": table_counter - 1,
                "paragraph_idx": previous_paragraph_idx,
            })

    return pd.DataFrame(rows)

# Figures
def extract_images(doc):
    """
    Parse all images from a .docx document using XML and return a DataFrame.
    """
    rows = []
    image_counter = 0

    for paragraph_idx, p in enumerate(doc.paragraphs):
        drawings = p._element.xpath(".//w:drawing")
        images_in_paragraph = len(drawings)

        for image_idx, drawing in enumerate(drawings):
            image_counter += 1

            # Extracting id of the image from the blip element
            blips = drawing.xpath(".//a:blip")
            r_id = blips[0].get(qn("r:embed")) if blips else None
            # Extracting the image
            image_bytes = doc.part.related_parts[r_id].blob

            rows.append({
                "image_no": image_counter,
                "image_idx_in_paragraph": image_idx,
                "images_in_paragraph": images_in_paragraph,
                "paragraph_idx": paragraph_idx,
                "r_id": r_id,
                "image_bytes": image_bytes,
            })

    return pd.DataFrame(rows)

def clean_text(text):
    """
    Clean the text by removing extra whitespace and normalizing it.
    """
    return re.sub(r"\s+", " ", text).strip().lower()

# List of Tables
LOT_TITLES = {
    "spis tabel": "Polish",
    "lista tabel": "Polish",
    "list of tables": "English",
}
LOF_TITLES = {
    "spis rysunków": "Polish",
    "lista rysunków": "Polish",
    "list of figures": "English",
    "table of figures": "English",
}
LIST_TITLES = LOT_TITLES.keys() | LOF_TITLES.keys()

def _detect_list(doc, titles, labels):
    for paragraph_idx, paragraph in enumerate(doc.paragraphs):
        title = clean_text(paragraph.text).rstrip(":")

        if title not in titles:
            continue

        language = titles[title]
        label_pattern = "|".join(re.escape(label) for label in labels[language])

        for entry in doc.paragraphs[paragraph_idx + 1:]:
            text = entry.text.strip()

            if not text:
                continue
            if clean_text(text).rstrip(":") in LIST_TITLES:
                return False, paragraph_idx, language, None

            match = re.match(
                rf"^\s*(?P<label>{label_pattern})\s*\.?\s*\d+",
                text,
                re.IGNORECASE,
            )
            if match:
                return True, paragraph_idx, language, match.group("label").rstrip(".")
            return False, paragraph_idx, language, None

        return False, paragraph_idx, language, None

    return False, None, None, None


def detect_list_of_tables(doc):
    labels = {
        "Polish": ("Tabela", "Tab"),
        "English": ("Table", "Tab"),
    }
    return _detect_list(doc, LOT_TITLES, labels)


# List of Figures
def detect_list_of_figures(doc):
    labels = {
        "Polish": ("Rysunek", "Rys"),
        "English": ("Figure", "Fig"),
    }
    return _detect_list(doc, LOF_TITLES, labels)

def _detect_list_content_range(doc, list_info, labels):
    """Return a half-open range containing a list heading and its contiguous entries."""
    start = list_info[1]

    if start is None:
        return None, None

    label_pattern = "|".join(re.escape(label) for label in labels)
    entry_regex = re.compile(
        rf"^\s*(?:{label_pattern})\s*\.?\s*\d+",
        re.IGNORECASE,
    )
    end = start + 1
    paragraphs = doc.paragraphs
    previous = paragraphs[start]._element

    for paragraph_idx in range(start + 1, len(paragraphs)):
        paragraph = paragraphs[paragraph_idx]
        element = paragraph._element
        # A list must not cross a table, drawing, embedded object or section break.
        if element.getprevious() is not previous or element.xpath(
            ".//w:drawing | .//w:pict | .//w:object | .//w:sectPr"
        ):
            break
        previous = element
        text = paragraph.text.strip()

        if not text:
            continue
        if not entry_regex.match(text):
            break

        end = paragraph_idx + 1

    return start, end


def detect_list_of_figures_content_range(doc):
    """Return the heading-and-content range of the existing figure list."""
    labels = ("Rysunek", "Rys", "Figure", "Fig")
    return _detect_list_content_range(doc, detect_list_of_figures(doc), labels)


def detect_list_of_tables_content_range(doc):
    """Return the heading-and-content range of the existing table list."""
    labels = ("Tabela", "Table", "Tab")
    return _detect_list_content_range(doc, detect_list_of_tables(doc), labels)
