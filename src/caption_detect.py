import re
import pandas as pd
from docx.oxml.ns import qn

DASHES = r"\-\u2013\u2014"
CAPTION_TEXT_TRIM_CHARS = " ([{,;:-\u2013\u2014"
SOURCE_TRIM_CHARS = " .;-\u2013\u2014"
FIGURE_NUMBER = r"\d+(?:\.\d+)*(?:[A-Za-z])?"
GROUP_SEPARATOR = r"(?:\b(?:i|oraz|and)\b|[,;&/]|\-|\u2013|\u2014)"
GROUP_ONE_ENTRY = "One entry for all images"
GROUP_SEPARATE_ENTRIES = "One entry per image"
GROUP_REVIEW_REQUIRED = "Review required"

def _clean_caption_text(text):
    if not text:
        return None

    return text.strip().rstrip(CAPTION_TEXT_TRIM_CHARS) or None


def _clean_source(source):
    if not source:
        return None

    clean_source = source.strip().strip(SOURCE_TRIM_CHARS)

    if (
        len(clean_source) >= 2
        and clean_source[0] == "("
        and clean_source[-1] == ")"
    ):
        clean_source = clean_source[1:-1].strip()

    if clean_source.endswith(")") and clean_source.count("(") < clean_source.count(")"):
        clean_source = clean_source[:-1].strip()

    return clean_source or None


def _plain_label(label):
    return re.sub(r"[^A-Za-z]", "", label).casefold()


def _caption_label_variants(caption_label):
    label = _plain_label(caption_label)

    if label.startswith("rys"):
        variants = {"rys", "rysunek"}
    elif label.startswith("fig"):
        variants = {"fig", "figure"}
    else:
        variants = {label}

    return sorted((variant for variant in variants if variant), key=len, reverse=True)


def _table_label_variants(table_label):
    label = _plain_label(table_label)
    variants = {"tab", "tabela", "table"} if label.startswith(("tab", "tabela")) else {label}
    return sorted((variant for variant in variants if variant), key=len, reverse=True)


def _source_label_variants(source_label):
    label = source_label.strip().casefold()
    label = re.sub(r"[:.;()\[\]{}]", "", label).strip()

    if "source" in label:
        variants = {"source"}
    elif "\u017ar\u00f3d" in label or "zrod" in label:
        variants = {"\u017ar\u00f3d\u0142o", "zrodlo"}
    else:
        variants = {label}

    return sorted((variant for variant in variants if variant), key=len, reverse=True)


def _number_separator_pattern(multi_number_separator=None):
    separators = {",", ";", "&", "/", r"\u2013", r"\u2014", "-", "and", "i", "oraz"}

    if multi_number_separator:
        separators.add(re.escape(multi_number_separator))

    words = [separator for separator in separators if separator.isalpha()]
    symbols = [separator for separator in separators if not separator.isalpha()]
    parts = []

    if words:
        parts.append(rf"\b(?:{'|'.join(sorted(words))})\b")

    if symbols:
        parts.append(rf"(?:{'|'.join(sorted(symbols))})")

    return "|".join(parts)


def _caption_regex(caption_label, allow_multiple_numbers=False, multi_number_separator=None):
    labels = "|".join(re.escape(label) for label in _caption_label_variants(caption_label))
    label_prefix = rf"(?:{labels})\s*\.?\s*(?:(?:nr|no|number)\s*\.?\s*)?"
    number = r"\d+(?:[.\-]\d+)*(?:[A-Za-z])?"
    number_ref = rf"(?:{label_prefix})?{number}"
    extra_numbers = ""

    if allow_multiple_numbers:
        separator = _number_separator_pattern(multi_number_separator)
        extra_numbers = rf"(?:\s*(?:{separator})\s*{number_ref})*"

    return re.compile(
        rf"""
        ^\s*
        (?P<label>
            {label_prefix}
            (?P<number>{number})?
            {extra_numbers}
        )
        \s*[\.\):;{DASHES}]?\s*
        (?P<body>.*)
        $
        """,
        re.IGNORECASE | re.VERBOSE,
    )


def _refresh_figure_caption_regex(caption_label):
    """Match a figure label and all explicitly connected numbers before caption text."""
    labels = "|".join(re.escape(label) for label in _caption_label_variants(caption_label))
    label_prefix = rf"(?:{labels})\s*\.?\s*(?:(?:nr|no|number)\s*\.?\s*)?"
    extra_number = rf"(?:\s*{GROUP_SEPARATOR}\s*(?:{label_prefix})?{FIGURE_NUMBER})*"
    return re.compile(
        rf"""
        ^\s*
        (?P<label>{label_prefix}(?P<number>{FIGURE_NUMBER}){extra_number})
        \s*[\.\):;]?\s*
        (?P<body>.*)
        $
        """,
        re.IGNORECASE | re.VERBOSE,
    )


def _table_caption_regex(table_label):
    labels = "|".join(re.escape(label) for label in _table_label_variants(table_label))
    label_prefix = rf"(?:{labels})\s*\.?\s*(?:(?:nr|no|number)\s*\.?\s*)?"
    number = r"\d+(?:[.\-]\d+)*(?:[A-Za-z])?"
    return re.compile(
        rf"""
        ^\s*
        (?P<label>{label_prefix}(?P<number>{number})?)
        \s*[\.\):;{DASHES}]?\s*
        (?P<body>.*)
        $
        """,
        re.IGNORECASE | re.VERBOSE,
    )


def _source_regex(source_label):
    sources = "|".join(
        re.escape(label) for label in _source_label_variants(source_label)
    )

    return re.compile(
        rf"""
        (?P<open>[\(\[\{{]?)\s*
        (?<!\w)
        (?:{sources})
        (?!\w)
        \s*[:;{DASHES}]?\s*
        """,
        re.IGNORECASE | re.VERBOSE,
    )


def _word_citation_regex():
    return re.compile(
        r"""
        (?P<citation>
            \s*
            [\(\[]\s*
            [^()\[\]]+
            (?:19|20)\d{2}[a-z]?
            [^()\[\]]*
            [\)\]]
            \s*$
        )
        """,
        re.IGNORECASE | re.VERBOSE,
    )


def _split_text_and_source(text, source_label):
    if not text:
        return None, None

    source_match = _source_regex(source_label).search(text)

    if source_match:
        caption_text = _clean_caption_text(text[:source_match.start()])
        source = _clean_source(text[source_match.end():])

        return caption_text, source

    citation_match = _word_citation_regex().search(text)

    if citation_match:
        caption_text = _clean_caption_text(text[:citation_match.start()])
        source = _clean_source(citation_match.group("citation"))

        return caption_text, source

    return _clean_caption_text(text), None


def _label_status(match):
    if not match:
        return "missing_label"

    if not match.group("number"):
        return "missing_number"

    return "found_label"


def _paragraph_visible_text(paragraph):
    """
    Return visible paragraph text from DOCX XML.

    Word citations are stored as fields: w:instrText contains the field command
    such as CITATION, while the rendered citation text is stored in w:t nodes.
    """

    texts = [
        node.text or ""
        for node in paragraph._element.xpath(".//w:t")
    ]

    return "".join(texts).strip()


def split_caption(
    caption,
    caption_label,
    source_label,
    allow_multiple_numbers=False,
    multi_number_separator=None,
):
    """
    Split a raw caption into label, text, source and label status.
    """

    if not caption:
        return None, None, None, "missing_label"

    match = _caption_regex(
        caption_label,
        allow_multiple_numbers=allow_multiple_numbers,
        multi_number_separator=multi_number_separator,
    ).match(caption.strip())
    status = _label_status(match)

    if not match:
        text, source = _split_text_and_source(caption.strip(), source_label)
        return None, text, source, status

    label = re.sub(r"\s+", " ", match.group("label")).strip()
    body = match.group("body").strip()
    text, source = _split_text_and_source(body, source_label)

    return label, text, source, status


def split_table_caption(caption, table_label, source_label):
    """Split a table caption into label, text, source and label status."""
    if not caption:
        return None, None, None, "missing_label"

    match = _table_caption_regex(table_label).match(caption.strip())
    status = _label_status(match)

    if not match:
        text, source = _split_text_and_source(caption.strip(), source_label)
        return None, text, source, status

    label = re.sub(r"\s+", " ", match.group("label")).strip()
    text, source = _split_text_and_source(match.group("body").strip(), source_label)
    return label, text, source, status


def _next_caption(doc, paragraph_idx, caption_regex, used_indices, limit=3):
    """Return the first matching caption from the next non-empty paragraphs."""
    checked = 0

    for candidate_idx in range(paragraph_idx + 1, len(doc.paragraphs)):
        text = _paragraph_visible_text(doc.paragraphs[candidate_idx])

        if not text:
            continue

        checked += 1
        match = None if candidate_idx in used_indices else caption_regex.match(text)

        if match:
            used_indices.add(candidate_idx)
            return match
        if checked == limit:
            break

    return None


def _classify_figure_group(label, numbers, images_count):
    if images_count == 1:
        return "Single"
    if len(numbers) == 1:
        return GROUP_ONE_ENTRY

    # Expand an explicit integer range only when it maps exactly to the image count.
    if len(numbers) == 2 and re.search(r"[\-\u2013\u2014]", label):
        if numbers[0].isdigit() and numbers[1].isdigit():
            start, end = map(int, numbers)
            if end - start + 1 == images_count:
                return GROUP_SEPARATE_ENTRIES
            return GROUP_REVIEW_REQUIRED

    if len(numbers) == images_count:
        return GROUP_SEPARATE_ENTRIES

    return GROUP_REVIEW_REQUIRED


def refresh_figures_caption_detect(doc, img_info, config):
    """Find figure captions and classify each image group independently."""
    columns = [
        "image_no", "image_idx_in_paragraph", "images_in_paragraph", "paragraph_idx",
        "image_bytes", "page", "text", "detected_caption", "detected_numbers",
        "suggested_treatment",
    ]

    if img_info.empty:
        return pd.DataFrame(columns=columns)

    rows = []
    ordered = img_info.sort_values("image_no")
    used_caption_indices = set()
    caption_regex = _refresh_figure_caption_regex(config["caption_label"])

    for _, group in ordered.groupby("paragraph_idx", sort=False):
        group = group.sort_values("image_idx_in_paragraph")
        first_image = group[group["image_idx_in_paragraph"] == 0]

        if first_image.empty:
            continue

        match = _next_caption(
            doc,
            int(first_image.iloc[0]["paragraph_idx"]),
            caption_regex,
            used_caption_indices,
        )

        if not match:
            continue

        detected_label = re.sub(r"\s+", " ", match.group("label")).strip()
        numbers = re.findall(FIGURE_NUMBER, detected_label)
        treatment = _classify_figure_group(detected_label, numbers, len(group))
        text, _ = _split_text_and_source(match.group("body").strip(), config["source_label"])

        for _, image in group.iterrows():
            rows.append({
                "image_no": image["image_no"],
                "image_idx_in_paragraph": image["image_idx_in_paragraph"],
                "images_in_paragraph": image["images_in_paragraph"],
                "paragraph_idx": image["paragraph_idx"],
                "image_bytes": image["image_bytes"],
                "page": image.get("page"),
                "text": text,
                "detected_caption": detected_label,
                "detected_numbers": ", ".join(numbers),
                "suggested_treatment": treatment,
            })

    return pd.DataFrame(rows, columns=columns)

def find_image_captions(
    doc,
    df_images,
    caption_label,
    source_label,
    max_lookahead=5,
    multi_image_strategy="composite",
    multi_number_separator=None,
):
    """
    Finds raw captions below images and splits them into label, text and source.

    Parameters:
    - doc: Document object from python-docx
    - df_images: DataFrame returned by extract_images()
    - caption_label: caption prefix provided by user, e.g. "Rys.", "Figure"
    - source_label: source prefix in captions, e.g. "Zrodlo:", "Source:"
    - max_lookahead: how many paragraphs below image should be checked

    Returns:
    - DataFrame with images and detected captions
    """

    columns = list(dict.fromkeys([
        *df_images.columns, "image_nos", "caption_paragraph_idx", "caption",
        "label", "label_status", "text", "source", "caption_status",
    ]))
    if df_images.empty:
        return pd.DataFrame(columns=columns)

    rows = []
    grouped_paragraphs = set()
    # Read all existing numbers regardless of the selected output strategy.
    caption_regex = _caption_regex(
        caption_label, allow_multiple_numbers=True,
        multi_number_separator=multi_number_separator,
    )

    for _, img in df_images.iterrows():
        image_paragraph_idx = int(img["paragraph_idx"])
        images_in_paragraph = int(img.get("images_in_paragraph", 1))
        is_multi_image_group = images_in_paragraph > 1

        if is_multi_image_group and image_paragraph_idx in grouped_paragraphs:
            continue

        if is_multi_image_group:
            grouped_paragraphs.add(image_paragraph_idx)
            group_images = df_images[
                df_images["paragraph_idx"] == image_paragraph_idx
            ]
            image_row = img.to_dict()
            image_nos = group_images["image_no"].astype(int).tolist()
            image_row["image_no"] = ", ".join(map(str, image_nos))
            image_row["image_nos"] = image_nos
            image_row["image_bytes"] = group_images["image_bytes"].tolist()
            if image_row.get("group_strategy") not in {"composite", "separate"}:
                image_row["group_strategy"] = multi_image_strategy
        else:
            image_row = img.to_dict()
            image_row["image_nos"] = [int(img["image_no"])]
            image_row["group_strategy"] = None

        caption_text = None
        caption_paragraph_idx = None
        caption_status = "not_found"

        for j in range(
            image_paragraph_idx + 1,
            min(image_paragraph_idx + 1 + max_lookahead, len(doc.paragraphs)),
        ):
            text = _paragraph_visible_text(doc.paragraphs[j])

            if not text:
                continue

            if caption_regex.match(text):
                caption_text = text
                caption_paragraph_idx = j
                caption_status = "found_by_label"
                break

            caption_text = text
            caption_paragraph_idx = j
            caption_status = "candidate_no_label"
            break

        label, clean_text, source, label_status = split_caption(
            caption_text,
            caption_label=caption_label,
            source_label=source_label,
            allow_multiple_numbers=True,
            multi_number_separator=multi_number_separator,
        )

        rows.append({
            **image_row,
            "caption_paragraph_idx": caption_paragraph_idx,
            "caption": caption_text,
            "label": label,
            "label_status": label_status,
            "text": clean_text,
            "source": source,
            "caption_status": caption_status,
        })

    return pd.DataFrame(rows, columns=columns)


def _table_boundaries(doc):
    """Map body tables to the nearest paragraph indices before and after them."""
    rows = []
    paragraph_idx = 0
    previous_paragraph_idx = None
    body_children = list(doc.element.body.iterchildren())

    for block_idx, child in enumerate(body_children):
        if child.tag == qn("w:p"):
            previous_paragraph_idx = paragraph_idx
            paragraph_idx += 1
            continue
        if child.tag != qn("w:tbl"):
            continue

        after_idx = next(
            (paragraph_idx for next_child in body_children[block_idx + 1:] if next_child.tag == qn("w:p")),
            None,
        )
        rows.append({
            "table_no": len(rows) + 1,
            "block_idx": block_idx,
            "before_paragraph_idx": previous_paragraph_idx,
            "after_paragraph_idx": after_idx,
        })

    return pd.DataFrame(rows)


def _tables_with_boundaries(doc, table_info):
    boundary_columns = ["table_no", "block_idx", "before_paragraph_idx", "after_paragraph_idx"]
    return table_info.merge(_table_boundaries(doc)[boundary_columns], on="table_no", how="left")


def _nearest_table_caption(doc, start_idx, direction, caption_regex, used_indices, limit=3):
    checked = 0
    paragraph_idx = start_idx

    while paragraph_idx is not None and 0 <= paragraph_idx < len(doc.paragraphs):
        current_idx = paragraph_idx
        text = _paragraph_visible_text(doc.paragraphs[current_idx])
        paragraph_idx += direction
        if not text:
            continue

        checked += 1
        match = None if current_idx in used_indices else caption_regex.match(text)
        if match:
            used_indices.add(current_idx)
            return match
        if checked == limit:
            break

    return None


def refresh_tables_caption_detect(doc, table_info, config):
    """Find existing table captions and prepare their content for list refresh."""
    columns = ["table_no", "table_idx", "page", "text"]

    if table_info.empty:
        return pd.DataFrame(columns=columns)

    tables = _tables_with_boundaries(doc, table_info)
    caption_regex = _table_caption_regex(config["table_caption_label"])
    caption_above = config["table_caption_position"] in {"caption_above", "label_above_source_below"}
    rows = []
    used_caption_indices = set()

    for _, table in tables.sort_values("table_no").iterrows():
        boundary = table["before_paragraph_idx"] if caption_above else table["after_paragraph_idx"]
        start_idx = None if pd.isna(boundary) else int(boundary)
        match = _nearest_table_caption(
            doc,
            start_idx,
            -1 if caption_above else 1,
            caption_regex,
            used_caption_indices,
        )

        if not match or not match.group("number"):
            continue

        text, _ = _split_text_and_source(match.group("body").strip(), config["source_label"])
        rows.append({
            "table_no": table["table_no"],
            "table_idx": table["table_idx"],
            "page": table.get("page"),
            "text": text,
        })

    return pd.DataFrame(rows, columns=columns)


def _scan_table_paragraphs(doc, paragraph_indices, table_label):
    caption_regex = _table_caption_regex(table_label)
    first_text = first_idx = None

    for paragraph_idx in paragraph_indices:
        if paragraph_idx is None or not 0 <= paragraph_idx < len(doc.paragraphs):
            continue

        text = _paragraph_visible_text(doc.paragraphs[paragraph_idx])
        if not text:
            continue
        if first_text is None:
            first_text, first_idx = text, paragraph_idx
        if caption_regex.match(text):
            return text, paragraph_idx, "found_by_label"

    if first_text is not None:
        return first_text, first_idx, "candidate_no_label"
    return None, None, "not_found"


def _source_from_below(doc, start_idx, source_label, max_lookaround):
    if start_idx is None:
        return None, None

    source_regex = _source_regex(source_label)
    for paragraph_idx in range(start_idx, min(start_idx + max_lookaround, len(doc.paragraphs))):
        text = _paragraph_visible_text(doc.paragraphs[paragraph_idx])
        if not text:
            continue

        source_match = source_regex.search(text)
        if source_match:
            return _clean_source(text[source_match.end():]), paragraph_idx

    return None, None


def find_table_captions(
    doc,
    df_tables,
    table_label,
    source_label,
    table_caption_position="caption_above",
    max_lookaround=5,
):
    """Find and parse table captions according to their configured position."""
    columns = list(dict.fromkeys([
        *df_tables.columns, "block_idx", "before_paragraph_idx", "after_paragraph_idx",
        "caption_paragraph_idx", "source_paragraph_idx", "caption", "label",
        "label_status", "text", "source", "caption_status", "table_caption_position",
    ]))
    if df_tables.empty:
        return pd.DataFrame(columns=columns)

    rows = []

    for _, table in _tables_with_boundaries(doc, df_tables).iterrows():
        before_idx, after_idx = table.get("before_paragraph_idx"), table.get("after_paragraph_idx")
        caption_above = table_caption_position in {"caption_above", "label_above_source_below"}
        boundary = before_idx if caption_above else after_idx
        start = None if pd.isna(boundary) else int(boundary)

        if start is None:
            paragraph_indices = []
        elif caption_above:
            paragraph_indices = range(start, max(start - max_lookaround, -1), -1)
        else:
            paragraph_indices = range(start, min(start + max_lookaround, len(doc.paragraphs)))

        caption, caption_idx, caption_status = _scan_table_paragraphs(
            doc, paragraph_indices, table_label
        )
        label, text, source, label_status = split_table_caption(
            caption, table_label, source_label
        )
        source_idx = caption_idx

        if table_caption_position == "label_above_source_below":
            source, source_idx = _source_from_below(
                doc,
                None if pd.isna(after_idx) else int(after_idx),
                source_label,
                max_lookaround,
            )

        rows.append({
            **table.to_dict(),
            "caption_paragraph_idx": caption_idx,
            "source_paragraph_idx": source_idx,
            "caption": caption,
            "label": label,
            "label_status": label_status,
            "text": text,
            "source": source,
            "caption_status": caption_status,
            "table_caption_position": table_caption_position,
        })

    return pd.DataFrame(rows, columns=columns)
