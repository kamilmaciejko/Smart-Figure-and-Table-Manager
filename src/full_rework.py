"""Streamlit workflow for rebuilding captions and Word caption lists."""

from io import BytesIO

import streamlit as st
from docx import Document

from src.ai_config import improve_list_rows
from src.caption_detect import find_image_captions, find_table_captions
from src.docx_edit import docx_to_bytes, edit_docx_captions, replace_caption_lists
from src.ui_helpers import (
    READ_ONLY_CAPTION_COLUMNS,
    READ_ONLY_TABLE_COLUMNS,
    apply_group_strategies,
    build_changed_captions,
    build_changed_table_captions,
    build_figure_list_rows,
    build_table_list_rows,
    caption_display_df,
    clean_caption_cell,
    clean_page,
    filter_rework_objects,
    highlight_ai_rows,
    highlight_attention_rows,
    image_preview,
    merge_list_edits,
    merge_caption_edits,
    object_removal_editor,
    output_file_name,
    render_changed_captions,
    table_preview,
    table_caption_display_df,
)
from src.word_pages import WordPageDetectionError, detect_object_pages


def _reset_results(stage="object_selection"):
    st.session_state.caption_stage = stage
    for key in (
        "accepted_captions", "accepted_table_captions", "edited_docx_bytes",
        "edited_docx_name", "caption_edited_docx_bytes", "table_of_figures", "list_of_tables",
        "full_ai_version",
    ):
        st.session_state[key] = None


def _object_selection(doc, images, tables, file_signature):
    st.header("Full Rework: delete invalid objects")
    st.caption("Step 2 of 6")
    st.caption("Exclude decorative objects and correct how individual image groups are treated.")
    excluded_images = st.session_state.excluded_image_nos
    excluded_tables = st.session_state.excluded_table_nos
    pending = st.session_state.pending_object_removal

    if pending:
        image_text = ", ".join(map(str, pending["images"])) or "none"
        table_text = ", ".join(map(str, pending["tables"])) or "none"
        st.error(f"Confirm removal from this rework. Images: {image_text}; tables: {table_text}.")
        confirm_col, cancel_col = st.columns(2)
        with confirm_col:
            if st.button("Confirm removal", type="primary"):
                st.session_state.excluded_image_nos = sorted(
                    set(excluded_images) | set(pending["images"])
                )
                st.session_state.excluded_table_nos = sorted(
                    set(excluded_tables) | set(pending["tables"])
                )
                st.session_state.pending_object_removal = None
                st.session_state.object_selection_version += 1
                _reset_results()
                st.rerun()
        with cancel_col:
            if st.button("Cancel removal"):
                st.session_state.pending_object_removal = None
                st.rerun()
        st.stop()

    st.write(f"Included: {len(images)} image(s), {len(tables)} table(s).")
    key = f"{file_signature[1][:10]}_{st.session_state.object_selection_version}"
    remove_images, remove_tables, strategies = object_removal_editor(images, tables, doc, key)

    # A group has one strategy. Rebuild the editor after either row is changed
    # so every image in that paragraph immediately shows the same selection.
    current_strategies = (
        images.loc[images["images_in_paragraph"].gt(1)]
        .drop_duplicates("paragraph_idx")
        .set_index("paragraph_idx")["group_strategy"]
        .to_dict()
    )
    changed_strategies = {
        paragraph_idx: strategy for paragraph_idx, strategy in strategies.items()
        if current_strategies.get(paragraph_idx) != strategy
    }
    if changed_strategies:
        st.session_state.group_strategies.update(changed_strategies)
        st.session_state.object_selection_version += 1
        _reset_results()
        st.rerun()

    if st.button(
        "Remove selected objects", type="primary",
        disabled=not remove_images and not remove_tables,
    ):
        st.session_state.pending_object_removal = {
            "images": remove_images, "tables": remove_tables,
        }
        st.rerun()

    if excluded_images or excluded_tables:
        st.caption(f"Excluded: {len(excluded_images)} image(s), {len(excluded_tables)} table(s).")
        if st.button("Restore all excluded objects"):
            st.session_state.excluded_image_nos = []
            st.session_state.excluded_table_nos = []
            st.session_state.object_selection_version += 1
            _reset_results()
            st.rerun()

    back_col, continue_col = st.columns(2)
    with back_col:
        if st.button("Back to caption format"):
            st.session_state.approved_caption_settings = None
            st.session_state.group_strategies = {}
            _reset_results()
            st.rerun()
    with continue_col:
        if st.button(
            "Continue to caption editing", type="primary",
            disabled=images.empty and tables.empty,
        ):
            _reset_results(stage="edit_table")
            st.rerun()


def _caption_frames(doc, images, tables, settings):
    figures = find_image_captions(
        doc, images, settings["caption_label"], settings["source_label"],
        multi_image_strategy=settings["multi_image_strategy"],
        multi_number_separator=settings["multi_number_separator"],
    )
    table_captions = find_table_captions(
        doc, tables, settings["table_caption_label"], settings["source_label"],
        settings["table_caption_position"],
    )
    return figures, table_captions


def _caption_editor(figures, tables, file_signature, doc):
    st.header("Full Rework: edit captions")
    st.caption("Step 3 of 6")
    st.caption(
        "Review the detected text and source. Highlighted rows need attention; "
        "target labels and numbering are added in the final preview."
    )
    version = st.session_state.object_selection_version
    key = f"{file_signature[1][:10]}_{version}"
    accepted_figures = st.session_state.accepted_captions
    accepted_tables = st.session_state.accepted_table_captions

    figure_base = accepted_figures if accepted_figures is not None else figures
    edited_figures = caption_display_df(figure_base)
    if not figures.empty:
        st.markdown("### Figure captions")
        edited_figures = st.data_editor(
            edited_figures.style.apply(highlight_attention_rows, axis=1),
            width="stretch", hide_index=True, disabled=READ_ONLY_CAPTION_COLUMNS,
            column_config={
                "image_no": st.column_config.NumberColumn("No."),
                "page": st.column_config.NumberColumn("Page"),
                "text": st.column_config.TextColumn("Text"),
                "source": st.column_config.TextColumn("Source"),
                "image_preview": st.column_config.ImageColumn("Preview", width="small"),
                "caption_status": st.column_config.TextColumn("Caption status"),
                "label_status": st.column_config.TextColumn("Label status"),
            },
            key=f"caption_editor_{key}",
        )

    table_base = (accepted_tables if accepted_tables is not None else tables).copy()
    if not table_base.empty:
        table_base["table_preview"] = table_base["table_idx"].apply(
            lambda index: table_preview(doc, index)
        )
    edited_tables = table_caption_display_df(table_base)
    if not tables.empty:
        st.markdown("### Table captions")
        edited_tables = st.data_editor(
            edited_tables.style.apply(highlight_attention_rows, axis=1),
            width="stretch", hide_index=True, disabled=READ_ONLY_TABLE_COLUMNS,
            column_config={
                "table_no": st.column_config.NumberColumn("No."),
                "page": st.column_config.NumberColumn("Page"),
                "text": st.column_config.TextColumn("Text"),
                "source": st.column_config.TextColumn("Source"),
                "table_preview": st.column_config.TextColumn("Preview"),
                "caption_status": st.column_config.TextColumn("Caption status"),
                "label_status": st.column_config.TextColumn("Label status"),
            },
            key=f"table_caption_editor_{key}",
        )

    captions_to_create = sum(
        int((frame["caption_status"] == "not_found").sum())
        for frame in (figures, tables) if not frame.empty
    )
    missing_text = any(
        not frame.empty and frame["text"].apply(clean_caption_cell).eq("").any()
        for frame in (edited_figures, edited_tables)
    )
    if captions_to_create:
        st.info(
            f"{captions_to_create} missing caption(s) will be created at their object positions."
        )
    if missing_text:
        st.error("Every included object must have caption text.")

    back_col, accept_col = st.columns(2)
    with back_col:
        if st.button("Back to object selection"):
            _reset_results(stage="object_selection")
            st.rerun()
    with accept_col:
        if st.button("Accept all captions", type="primary", disabled=missing_text):
            st.session_state.accepted_captions = merge_caption_edits(figures, edited_figures)
            st.session_state.accepted_table_captions = merge_caption_edits(tables, edited_tables)
            st.session_state.caption_stage = "review_table"
            st.session_state.edited_docx_bytes = None
            st.session_state.caption_edited_docx_bytes = None
            st.session_state.table_of_figures = None
            st.session_state.list_of_tables = None
            st.rerun()


def _review_captions(uploaded_bytes, file_name, settings):
    figures = st.session_state.accepted_captions
    tables = st.session_state.accepted_table_captions
    st.header("Full Rework: review changes")
    st.caption("Step 4 of 6")
    st.info(f"Accepted captions: {len(figures) + len(tables)}")
    changed_figures = build_changed_captions(figures, settings)
    changed_tables = build_changed_table_captions(tables, settings)

    if changed_figures:
        st.subheader("Figure captions changed")
        render_changed_captions(changed_figures, settings)
    if changed_tables:
        st.subheader("Table captions changed")
        render_changed_captions(changed_tables, settings)

    back_col, confirm_col = st.columns(2)
    with back_col:
        if st.button("Back and edit"):
            st.session_state.caption_stage = "edit_table"
            st.rerun()
    with confirm_col:
        if st.button("Apply captions to DOCX", type="primary"):
            all_changes = [*changed_figures, *changed_tables]
            edited_doc = edit_docx_captions(
                Document(BytesIO(uploaded_bytes)), all_changes, settings
            )
            st.session_state.caption_edited_docx_bytes = docx_to_bytes(edited_doc)
            st.session_state.edited_docx_bytes = None
            st.session_state.edited_docx_name = output_file_name(file_name)
            st.session_state.caption_stage = "list_choice"
            st.rerun()


def _choose_lists(settings):
    st.header("Full Rework: create lists")
    st.caption("Step 5 of 6")
    st.caption("Existing lists will be replaced; missing lists will be added to the document.")
    back_col, skip_col, add_col = st.columns(3)
    with back_col:
        if st.button("Back to captions"):
            st.session_state.caption_stage = "edit_table"
            st.session_state.caption_edited_docx_bytes = None
            st.rerun()
    with skip_col:
        if st.button("Download without lists"):
            st.session_state.edited_docx_bytes = st.session_state.caption_edited_docx_bytes
            st.session_state.caption_stage = "docx_ready"
            st.rerun()
    with add_col:
        if st.button("Create lists", type="primary"):
            st.session_state.table_of_figures = build_figure_list_rows(
                st.session_state.accepted_captions, settings,
                st.session_state.get("image_pages"),
            )
            st.session_state.list_of_tables = build_table_list_rows(
                st.session_state.accepted_table_captions, settings,
                st.session_state.get("table_pages"),
            )
            st.session_state.caption_stage = "list_edit"
            st.rerun()


def _list_editor(rows, key):
    """Render list rows while keeping internal AI grouping columns hidden."""
    ai_result = "decision" in rows.columns
    previews = [column for column in ("image_preview", "table_preview") if column in rows]
    columns = ["label", "text", "page"]
    if ai_result:
        columns = ["label", "original_text", "text", "decision", "edit_type", "page"]
    columns += previews
    display = rows[columns]
    if ai_result:
        display = display.style.apply(highlight_ai_rows, axis=1)
    disabled = ["label", "page", *previews]
    if ai_result:
        disabled += ["original_text", "decision", "edit_type"]
    edited = st.data_editor(
        display, width="stretch", hide_index=True, disabled=disabled,
        column_config={
            "label": st.column_config.TextColumn("Label"),
            "original_text": st.column_config.TextColumn("Original text"),
            "text": st.column_config.TextColumn("Text"),
            "page": st.column_config.NumberColumn("Page"),
            "image_preview": st.column_config.ImageColumn("Image preview", width="small"),
            "table_preview": st.column_config.TextColumn("Table preview"),
            "decision": st.column_config.TextColumn("AI decision"),
            "edit_type": st.column_config.TextColumn("Edit type"),
        },
        key=key,
    )
    return merge_list_edits(rows, edited)


def _edit_lists(
    settings, caption_bytes, has_figure_heading, has_table_heading,
    image_count, table_count,
):
    st.header(
        "Listy tabel i rysunków" if settings["language"] == "Polish"
        else "Lists of Tables and Figures"
    )
    st.caption("Step 5 of 6")
    st.caption("Edit descriptions if needed. Page numbers are recalculated when the DOCX is created.")
    tables = st.session_state.list_of_tables
    figures = st.session_state.table_of_figures
    ai_version = st.session_state.get("full_ai_version") or 0
    caption_doc = Document(BytesIO(caption_bytes))

    if not figures.empty and "image_preview" not in figures and "image_bytes" in figures:
        figures = figures.copy()
        figures["image_preview"] = figures["image_bytes"].apply(image_preview)
    if not tables.empty and "table_preview" not in tables and "table_idx" in tables:
        tables = tables.copy()
        tables["table_preview"] = tables["table_idx"].apply(
            lambda index: table_preview(caption_doc, index)
        )

    edited_tables = tables
    if not tables.empty:
        st.markdown("### Lista tabel" if settings["language"] == "Polish" else "### List of Tables")
        edited_tables = _list_editor(tables, f"list_of_tables_editor_{ai_version}")

    edited_figures = figures
    if not figures.empty:
        st.markdown("### Spis rysunków" if settings["language"] == "Polish" else "### List of Figures")
        edited_figures = _list_editor(figures, f"table_of_figures_editor_{ai_version}")

    if st.session_state.pop("ai_review_requested", False):
        activation = st.session_state.get("llm_activation")
        config = st.session_state.get("llm_runtime_config")
        try:
            if not activation or activation.get("status") != "llm_active" or not config:
                st.error("AI functions not active.")
            else:
                with st.spinner("Reviewing list captions with AI..."):
                    ai_figures = improve_list_rows(
                        config, edited_figures, "figures", settings["language"]
                    )
                    ai_tables = improve_list_rows(
                        config, edited_tables, "tables", settings["language"]
                    )
                st.session_state.table_of_figures = ai_figures
                st.session_state.list_of_tables = ai_tables
                st.session_state.full_ai_version = ai_version + 1
                st.rerun()
        except Exception as error:
            st.warning(f"AI caption review failed: {error}")
        finally:
            st.session_state.ai_request_in_progress = False

    invalid = any(
        not frame.empty and (
            frame["text"].apply(clean_caption_cell).eq("").any()
            or frame["page"].apply(clean_page).eq("").any()
        )
        for frame in (edited_tables, edited_figures)
    )
    if invalid:
        st.error("Every list entry must contain text and a valid page number.")

    back_col, accept_col = st.columns(2)
    with back_col:
        if st.button("Back to caption review"):
            st.session_state.caption_stage = "review_table"
            st.session_state.table_of_figures = None
            st.session_state.list_of_tables = None
            st.rerun()
    with accept_col:
        if st.button("Create DOCX", type="primary", disabled=invalid):
            replace_figures = has_figure_heading or not edited_figures.empty
            replace_tables = has_table_heading or not edited_tables.empty
            try:
                # First establish the final list length and edited text, then paginate that layout.
                provisional_doc = replace_caption_lists(
                    Document(BytesIO(caption_bytes)), edited_figures, edited_tables, settings,
                    replace_figures=replace_figures, replace_tables=replace_tables,
                )
                provisional_bytes = docx_to_bytes(provisional_doc)
                with st.spinner("Updating final page numbers with Microsoft Word..."):
                    pages = detect_object_pages(provisional_bytes, image_count, table_count)

                final_figures = build_figure_list_rows(
                    st.session_state.accepted_captions, settings, pages["images"]
                )
                final_tables = build_table_list_rows(
                    st.session_state.accepted_table_captions, settings, pages["tables"]
                )
                final_figures["text"] = edited_figures["text"].tolist()
                final_tables["text"] = edited_tables["text"].tolist()
                final_doc = replace_caption_lists(
                    Document(BytesIO(provisional_bytes)), final_figures, final_tables, settings,
                    replace_figures=replace_figures, replace_tables=replace_tables,
                )
                st.session_state.table_of_figures = final_figures
                st.session_state.list_of_tables = final_tables
                st.session_state.edited_docx_bytes = docx_to_bytes(final_doc)
                st.session_state.caption_stage = "docx_ready"
                st.rerun()
            except WordPageDetectionError as error:
                st.error("Final page numbers could not be detected.")
                st.exception(error)


def run_full_rework(
    doc, uploaded_bytes, file_name, file_signature, df_images, df_tables,
    settings, has_figure_heading, has_table_heading,
):
    """Run the complete, stateful Full Rework workflow."""
    st.session_state.setdefault("group_strategies", {})
    stage = st.session_state.caption_stage
    images, tables = filter_rework_objects(
        df_images, df_tables,
        st.session_state.excluded_image_nos,
        st.session_state.excluded_table_nos,
    )
    images = apply_group_strategies(
        images, st.session_state.group_strategies, settings["multi_image_strategy"]
    )
    if stage == "object_selection":
        _object_selection(doc, images, tables, file_signature)
        return
    if images.empty and tables.empty:
        st.warning("No objects remain in the Full Rework selection.")
        if st.button("Back to object selection"):
            st.session_state.caption_stage = "object_selection"
            st.rerun()
        return

    if stage == "edit_table":
        figures, table_captions = _caption_frames(doc, images, tables, settings)
        if (not figures.empty and figures["page"].isna().any()) or (
            not table_captions.empty and table_captions["page"].isna().any()
        ):
            st.warning("Some object page numbers could not be detected.")
        _caption_editor(figures, table_captions, file_signature, doc)
    elif stage == "review_table":
        _review_captions(uploaded_bytes, file_name, settings)
    elif stage == "list_choice":
        _choose_lists(settings)
    elif stage == "list_edit":
        _edit_lists(
            settings, st.session_state.caption_edited_docx_bytes,
            has_figure_heading, has_table_heading, len(df_images), len(df_tables),
        )
    elif stage == "docx_ready":
        st.header("Full Rework: ready to download")
        st.caption("Step 6 of 6")
        st.success("The corrected DOCX file is ready.")
        st.download_button(
            "Download corrected DOCX", data=st.session_state.edited_docx_bytes,
            file_name=st.session_state.edited_docx_name,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary",
        )
        if st.button("Back to list options"):
            st.session_state.caption_stage = "list_choice"
            st.session_state.edited_docx_bytes = None
            st.rerun()
