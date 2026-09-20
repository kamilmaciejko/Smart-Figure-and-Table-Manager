"""Exercise both Streamlit workflows with a synthetic document and mocked Word pages."""

import hashlib
import unittest
from io import BytesIO
from unittest.mock import patch

import pandas as pd
from docx import Document
from streamlit.testing.v1 import AppTest

from src.docx_edit import docx_to_bytes
from test_docx_workflows import add_picture


class AppWorkflowTests(unittest.TestCase):
    def setUp(self):
        doc = Document()
        add_picture(doc.add_paragraph())
        doc.add_paragraph("Figure 8. Model performance (Source: study)")
        doc.add_paragraph("Table 4. Model scores (Source: study)")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "0.94"
        for text in ("List of Figures", "Figure 8. Model performance\t1", "List of Tables", "Table 4. Model scores\t1"):
            doc.add_paragraph(text)
        data = docx_to_bytes(doc)
        pages = {kind: pd.DataFrame([{"index": 1, "page": 1}]) for kind in ("images", "tables")}
        worker = patch("src.word_pages.detect_object_pages", return_value=pages)
        worker.start()
        self.addCleanup(worker.stop)
        self.app = AppTest.from_file("app.py", default_timeout=20).run()
        self.app.session_state["uploaded_docx_bytes"] = data
        self.app.session_state["uploaded_docx_name"] = "example.docx"
        self.app.session_state["uploaded_file_signature"] = ("example.docx", hashlib.sha256(data).hexdigest())
        self.app.run()
        self.assertFalse(self.app.exception)

    def click(self, label):
        next(button for button in self.app.button if button.label == label).click().run()
        self.assertFalse(self.app.exception)

    def test_full_rework_can_generate_captioned_document_and_lists(self):
        self.click("Full Tables and Figures Rework")
        self.app.selectbox(key="caption_language").select("English").run()
        self.click("Continue to object selection")
        session = self.app.session_state
        self.app = AppTest.from_file("app.py", default_timeout=20)
        self.app.session_state = session
        self.app.run()
        for label in ("Continue to caption editing", "Accept all captions", "Apply captions to DOCX", "Create lists", "Create DOCX"):
            self.click(label)
        self.assertEqual(self.app.session_state["caption_stage"], "docx_ready")
        result = Document(BytesIO(self.app.session_state["edited_docx_bytes"]))
        texts = [p.text for p in result.paragraphs]
        self.assertIn("Fig. 1. Model performance (Source: study)", texts)
        self.assertIn("Fig. 1. Model performance\t1", texts)
        self.assertEqual(len(result.inline_shapes), 1)
        self.assertEqual(result.tables[0].cell(0, 0).text, "0.94")

    def test_refresh_changes_lists_and_preserves_body_captions(self):
        self.click("Refresh List of Tables and Figures")
        self.click("Continue to list review")
        self.click("Accept refreshed lists")
        result = Document(BytesIO(self.app.session_state["refresh_docx_bytes"]))
        texts = [p.text for p in result.paragraphs]
        self.assertIn("Figure 8. Model performance (Source: study)", texts)
        self.assertIn("Figure 1. Model performance\t1", texts)
        self.assertIn("Table 1. Model scores\t1", texts)
        self.assertEqual(len(result.inline_shapes), 1)
        self.assertEqual(result.tables[0].cell(0, 0).text, "0.94")


if __name__ == "__main__":
    unittest.main()
