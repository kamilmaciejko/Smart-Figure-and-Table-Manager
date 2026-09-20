import unittest

import pandas as pd
from docx import Document

from src.ui_helpers import (
    READ_ONLY_TABLE_COLUMNS,
    TABLE_DISPLAY_COLUMNS,
    build_figure_list_rows,
    build_table_list_rows,
    table_caption_display_df,
    table_preview,
)


class ListPreviewTests(unittest.TestCase):
    def test_separate_figures_keep_matching_preview_bytes(self):
        captions = pd.DataFrame([{
            "image_no": "1, 2", "image_nos": [1, 2], "images_in_paragraph": 2,
            "image_bytes": [b"first", b"second"], "text": "Shared", "page": 3,
        }])
        settings = {"caption_label": "Fig.", "multi_image_strategy": "separate"}

        rows = build_figure_list_rows(captions, settings)

        self.assertEqual(rows["image_bytes"].tolist(), [b"first", b"second"])

    def test_table_rows_keep_word_table_index(self):
        captions = pd.DataFrame([{"table_no": 1, "table_idx": 0, "text": "Results", "page": 4}])
        settings = {"table_caption_label": "Table"}

        rows = build_table_list_rows(captions, settings)

        self.assertEqual(rows.loc[0, "table_idx"], 0)

    def test_table_preview_uses_table_content(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Model"
        table.cell(0, 1).text = "Score"

        self.assertEqual(table_preview(doc, 0), "Model | Score")

    def test_table_caption_editor_order_prioritizes_editing_fields(self):
        expected = [
            "table_preview", "label", "text", "source",
            "table_no", "page", "caption_status", "label_status",
        ]
        display = table_caption_display_df(pd.DataFrame([{"text": "Results"}]))

        self.assertEqual(TABLE_DISPLAY_COLUMNS, expected)
        self.assertEqual(display.columns.tolist(), expected)
        self.assertEqual(
            READ_ONLY_TABLE_COLUMNS,
            ["table_preview", "label", "table_no", "page", "caption_status", "label_status"],
        )


if __name__ == "__main__":
    unittest.main()
