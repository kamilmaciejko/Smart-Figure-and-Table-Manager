"""In-memory Word fixtures: no personal documents, Word installation or AI calls."""

import unittest
from io import BytesIO

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

from src.caption_detect import (
    GROUP_REVIEW_REQUIRED,
    GROUP_SEPARATE_ENTRIES,
    _classify_figure_group,
    find_image_captions,
    find_table_captions,
    refresh_figures_caption_detect,
    refresh_tables_caption_detect,
)
from src.docx_edit import docx_to_bytes, edit_docx_captions, replace_caption_lists
from src.object_detect import detect_list_of_figures, extract_images, extract_tables
from src.ui_helpers import build_changed_captions, build_changed_table_captions, filter_rework_objects


SETTINGS = {
    "language": "English", "caption_label": "Figure", "table_caption_label": "Table",
    "source_label": "Source:", "source_in_parentheses": True,
    "font_name": "Arial", "font_size": 10,
    "label_bold": True, "label_italic": False,
    "text_bold": False, "text_italic": False,
    "source_bold": False, "source_italic": True,
    "multi_image_strategy": "separate", "multi_number_separator": "–",
    "table_caption_position": "caption_above",
}


def add_picture(paragraph):
    buffer = BytesIO()
    Image.new("RGB", (12, 12), "blue").save(buffer, format="PNG")
    buffer.seek(0)
    paragraph.add_run().add_picture(buffer, width=Inches(0.5))


class DocxWorkflowTests(unittest.TestCase):
    def test_group_detection_and_exclusion_preserve_object_ids(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        add_picture(paragraph)
        add_picture(paragraph)
        doc.add_paragraph("Figure 7 and 8. Model results (Source: study)")
        images = extract_images(doc).assign(page=2)
        refreshed = refresh_figures_caption_detect(doc, images, SETTINGS)
        self.assertEqual(refreshed["suggested_treatment"].tolist(), [GROUP_SEPARATE_ENTRIES] * 2)
        selected, _ = filter_rework_objects(images, pd.DataFrame(), excluded_images=[1])
        captions = find_image_captions(doc, selected, "Figure", "Source:")
        self.assertEqual(captions.iloc[0]["image_nos"], [2])
        self.assertEqual(captions.iloc[0]["text"], "Model results")
        self.assertEqual(captions.iloc[0]["source"], "study")
        self.assertEqual(len(doc.inline_shapes), 2)

    def test_caption_edits_preserve_other_paragraphs_and_image_bytes(self):
        doc = Document()
        doc.add_paragraph("Introduction to the study.")
        add_picture(doc.add_paragraph())
        caption = doc.add_paragraph("Figure 9. Old results (Source: dataset)")
        caption.paragraph_format.keep_with_next = True
        doc.add_paragraph("Conclusion remains unchanged.")
        images = extract_images(doc).assign(page=1)
        captions = find_image_captions(doc, images, "Figure", "Source:")
        captions.loc[0, "text"] = "Updated results"
        changed = build_changed_captions(captions, SETTINGS)
        result = Document(BytesIO(docx_to_bytes(edit_docx_captions(doc, changed, SETTINGS))))
        self.assertEqual(result.paragraphs[0].text, "Introduction to the study.")
        self.assertEqual(result.paragraphs[-1].text, "Conclusion remains unchanged.")
        self.assertEqual(result.paragraphs[2].text, "Figure 1. Updated results (Source: dataset)")
        self.assertTrue(result.paragraphs[2].paragraph_format.keep_with_next)
        self.assertEqual(extract_images(result).iloc[0]["image_bytes"], images.iloc[0]["image_bytes"])

    def test_missing_captions_use_original_anchors_after_insertions(self):
        doc = Document()
        add_picture(doc.add_paragraph())
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Score"
        rows = [
            {"object_type": "figure", "object_paragraph_idx": 0, "caption_paragraph_idx": pd.NA,
             "caption_text": "Results", "new_number": "1", "source_text": pd.NA},
            {"object_type": "table", "table_idx": 0, "caption_text": "Scores", "new_number": "1",
             "caption_label": "Table", "source_text": "study", "table_caption_position": "label_above_source_below"},
        ]
        result = edit_docx_captions(doc, rows, SETTINGS)
        blocks = [element.tag for element in result.element.body if element.tag != qn("w:sectPr")]
        self.assertEqual(blocks, [qn("w:p")] * 3 + [qn("w:tbl"), qn("w:p")])
        self.assertEqual([p.text.strip() for p in result.paragraphs], ["", "Figure 1. Results", "Table 1. Scores", "(Source: study)"])
        self.assertEqual(result.tables[0].cell(0, 0).text, "Score")

    def test_all_table_caption_positions_round_trip(self):
        for position in ("caption_above", "caption_below", "label_above_source_below"):
            with self.subTest(position=position):
                doc = Document()
                text = "Table 8. Model scores"
                if position != "label_above_source_below":
                    text += " (Source: study)"
                if position != "caption_below":
                    doc.add_paragraph(text)
                doc.add_table(rows=1, cols=1).cell(0, 0).text = "0.94"
                if position == "caption_below":
                    doc.add_paragraph(text)
                elif position == "label_above_source_below":
                    doc.add_paragraph("Source: study")
                tables = extract_tables(doc).assign(page=3)
                settings = {**SETTINGS, "table_caption_position": position}
                captions = find_table_captions(doc, tables, "Table", "Source:", position)
                self.assertEqual(captions.iloc[0]["source"], "study")
                changed = build_changed_table_captions(captions, settings)
                result = edit_docx_captions(doc, changed, settings)
                refreshed = refresh_tables_caption_detect(result, extract_tables(result).assign(page=3), settings)
                self.assertEqual(refreshed["text"].tolist(), ["Model scores"])
                self.assertEqual(result.tables[0].cell(0, 0).text, "0.94")

    def test_lists_are_replaced_in_place_without_touching_body(self):
        doc = Document()
        for text in ("Preface", "List of Figures", "Figure 8. Old figure\t1", "List of Tables", "Table 7. Old table\t2", "Chapter one"):
            doc.add_paragraph(text)
        figures = [{"label": "Figure 1.", "text": "New figure", "page": 4}]
        tables = [{"label": "Table 1.", "text": "New table", "page": 5}]
        result = replace_caption_lists(doc, figures, tables, SETTINGS)
        self.assertEqual([p.text for p in result.paragraphs], [
            "Preface", "List of Figures", "Figure 1. New figure\t4",
            "List of Tables", "Table 1. New table\t5", "Chapter one",
        ])

    def test_list_replacement_stops_at_document_objects(self):
        for boundary in ("image", "table", "section"):
            with self.subTest(boundary=boundary):
                doc = Document()
                doc.add_paragraph("List of Figures")
                doc.add_paragraph("Figure 1. Old entry\t1")
                if boundary == "image":
                    add_picture(doc.add_paragraph())
                elif boundary == "table":
                    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Preserved"
                else:
                    doc.add_paragraph()._element.get_or_add_pPr().append(OxmlElement("w:sectPr"))
                doc.add_paragraph("Figure 2. Body caption")
                result = replace_caption_lists(doc, [{"label": "Figure 1.", "text": "New entry", "page": 1}], [], SETTINGS, replace_tables=False)
                self.assertEqual(result.paragraphs[-1].text, "Figure 2. Body caption")
                if boundary == "image":
                    self.assertEqual(len(result.inline_shapes), 1)
                elif boundary == "table":
                    self.assertEqual(result.tables[0].cell(0, 0).text, "Preserved")
                else:
                    self.assertEqual(len(result.sections), 2)

    def test_polish_list_headings_and_empty_lists(self):
        for title in ("Spis rysunków", "Lista rysunków"):
            doc = Document()
            doc.add_paragraph(title)
            self.assertEqual(detect_list_of_figures(doc), (False, 0, "Polish", None))
            doc.add_paragraph("Rys. 1. Wyniki\t2")
            self.assertEqual(detect_list_of_figures(doc), (True, 0, "Polish", "Rys"))

    def test_large_caption_range_requires_review_without_expansion(self):
        self.assertEqual(_classify_figure_group("Figure 1–999999999999", ["1", "999999999999"], 2), GROUP_REVIEW_REQUIRED)


if __name__ == "__main__":
    unittest.main()
