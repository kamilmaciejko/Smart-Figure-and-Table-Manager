"""Optional check against a real local Microsoft Word installation."""

import os
import unittest

from docx import Document

from src.docx_edit import docx_to_bytes
from src.word_pages import detect_object_pages
from test_docx_workflows import add_picture


@unittest.skipUnless(os.name == "nt" and os.environ.get("RUN_WORD_TESTS") == "1", "Set RUN_WORD_TESTS=1 on Windows with Microsoft Word installed.")
class WordIntegrationTests(unittest.TestCase):
    def test_physical_pages_in_real_word(self):
        doc = Document()
        add_picture(doc.add_paragraph())
        doc.add_paragraph("Figure 1. Test image")
        doc.add_page_break()
        doc.add_paragraph("Table 1. Test table")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "0.94"
        pages = detect_object_pages(docx_to_bytes(doc), expected_images=1, expected_tables=1)
        self.assertEqual(pages["images"]["page"].tolist(), [1])
        self.assertEqual(pages["tables"]["page"].tolist(), [2])


if __name__ == "__main__":
    unittest.main()
